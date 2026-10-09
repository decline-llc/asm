import ipaddress
import re
import shutil
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..pipeline import StageResult

HEADERS = {
    "总表": "序号/资产/端口/标题/技术栈/等级标注/来源/置信度/截图链接".split("/"),
    "equity": "cid,name,parent_cid,share_pct,level,website,email_suffix".split(","),
    "domains": "domain,registrable,tag5,confidence,wildcard,ip,sources".split(","),
    "ips": "ip,cid,ports_count,cluster_kind,cdn".split(","),
    "urls": "url,asset_id,title,status,len,tech,screenshot".split(","),
    "systems": "system_name,cid,tags,urls,vendor,version,source".split(","),
    "social": "name,role,company,email,phone,qq,platform,source".split(","),
    "todos": "type,target,evidence,severity,next_step".split(","),
}


def text_cell(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32700]
        # External titles, names and evidence never become executable spreadsheet formulas.
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def write_report(db, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    all_rows = {name: [] for name in HEADERS}
    companies = {r["cid"]: r for r in db.rows("SELECT * FROM companies")}
    assets = db.rows("SELECT * FROM assets ORDER BY host,port,proto")
    url_rows = db.rows("SELECT * FROM urls ORDER BY url")
    screenshots = {}
    for number, asset in enumerate(assets, 1):
        source = next((r["screenshot"] for r in url_rows if r["asset_id"] == asset["id"] and r["screenshot"]), "")
        link = ""
        if source and Path(source).is_file():
            dest = path.parent / "screenshots" / f"{number}.png"
            dest.parent.mkdir(parents=True, exist_ok=True)
            if Path(source).resolve() != dest.resolve():
                shutil.copyfile(source, dest)
            link = f"screenshots/{number}.png"
            screenshots[asset["id"]] = link
        tags = []
        host = asset["host"].lower()
        if any(k in host for k in ("uat", "test", "dev", "sit")):
            tags.append("UAT/测试环境")
        if "minio" in str(asset).lower():
            tags.append("MinIO对象存储")
        if any(k in host for k in ("hr.", "fin.", "erp.")):
            tags.append("HR/财务")
        if asset["cdn"]:
            tags.append("云部署")
        all_rows["总表"].append([number, asset["host"], asset["port"], asset["title"], asset["tech"],
                                    ",".join(tags), asset["source"], asset["confidence"], link])
    all_rows["equity"] = [[r.get(k) for k in HEADERS["equity"]] for r in companies.values()]
    for row in db.rows("SELECT * FROM domains ORDER BY domain"):
        row["ip"] = ",".join(r["value"] for r in db.rows(
            "SELECT DISTINCT value FROM dns_records WHERE domain=? AND rtype IN ('A','AAAA')", (row["domain"],)))
        all_rows["domains"].append([row.get(k) for k in HEADERS["domains"]])
    for row in db.rows("SELECT ip,min(cid) cid,count(DISTINCT port) ports_count,max(cdn) cdn "
                       "FROM assets WHERE ip IS NOT NULL AND ip!='' GROUP BY ip ORDER BY ip"):
        try:
            ip = ipaddress.ip_address(row["ip"])
            network = str(ipaddress.ip_network(f"{ip}/{'24' if ip.version == 4 else '64'}", strict=False))
            row["cluster_kind"] = db.get_fact("ip_clusters", {}).get(network, "其他")
        except ValueError:
            row["cluster_kind"] = "其他"
        all_rows["ips"].append([row.get(k) for k in HEADERS["ips"]])
    for row in url_rows:
        row["screenshot"] = screenshots.get(row["asset_id"], "")
        all_rows["urls"].append([row.get(k) for k in HEADERS["urls"]])
    all_rows["systems"] = [[r.get(k) for k in HEADERS["systems"]] for r in db.rows("SELECT * FROM systems")]
    for row in db.rows("SELECT * FROM people"):
        row["company"] = companies.get(row["cid"], {}).get("name", "")
        row["platform"] = row.get("weibo", "")
        all_rows["social"].append([row.get(k) for k in HEADERS["social"]])
    for row in db.rows("SELECT * FROM findings WHERE status!='rejected' ORDER BY id"):
        all_rows["todos"].append([row["kind"], row["url"], row["evidence"], row["severity"], row["status"]])
    for name, headers in HEADERS.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for values in all_rows[name]:
            sheet.append([text_cell(v) for v in values])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183153")
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F5F9")
                if isinstance(cell.value, str) and cell.value.startswith("screenshots/"):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        for col, header in enumerate(headers, 1):
            width = 48 if header in {"evidence", "资产", "url", "urls", "title", "标题", "screenshot", "截图链接"} else 24
            sheet.column_dimensions[get_column_letter(col)].width = width
        for row in sheet.iter_rows(min_row=2):
            lines = max((sum(max(1, math.ceil(sum(2 if ord(c)>255 else 1 for c in part) /
                            max(1, sheet.column_dimensions[cell.column_letter].width - 2)))
                            for part in str(cell.value or "").split("\n")) for cell in row), default=1)
            sheet.row_dimensions[row[0].row].height = min(300, max(22, lines * 15 + 6))
    book.save(path)
    book.close()
    return {name: len(rows) for name, rows in all_rows.items()}


def run(ctx):
    path = ctx.stage_dir.parent.parent / "report.xlsx"
    counts = write_report(ctx.db, path)
    ctx.write("report-counts.json", counts)
    return StageResult(count=counts["总表"], notes=[str(path)])
