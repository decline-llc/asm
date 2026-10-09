from dataclasses import replace

from openpyxl import load_workbook

from asm.conf import load_config
from asm.models import Asset, Domain
from asm.pipeline import run_pipeline
from asm.stages.s8_report import HEADERS, write_report


def test_report_contract_and_isolation(db, tmp_path):
    db.asset(Asset(host="api.example.invalid", ip="192.0.2.1", title="=1+1", source="fixture"))
    db.asset(Asset(host="fuzzy.example.invalid", title="quarantine", confidence="D", source="body"))
    db.domain(Domain("api.example.invalid"))
    path = tmp_path / "report.xlsx"
    counts = write_report(db, path)
    book = load_workbook(path, data_only=False)
    assert book.sheetnames == list(HEADERS)
    assert counts["总表"] == db.rows("SELECT COUNT(*) n FROM assets")[0]["n"] == 1
    for name, headers in HEADERS.items():
        assert [c.value for c in book[name][1]] == headers
        assert book[name].max_row == counts[name] + 1
    assert book["总表"]["D2"].data_type == "s" and book["总表"]["D2"].value == "'=1+1"
    assert "fuzzy" not in str(list(book["总表"].values))
    book.close()


def test_resume_exact_config(project, tmp_path, monkeypatch):
    config = load_config("demo", project)
    config = replace(config, root=tmp_path, data={**config.data, "output": {"dir": "data"}})
    monkeypatch.setattr("asm.pipeline.load_config", lambda *_: config)
    first = run_pipeline(stages=["1", "3"], passive_only=True)
    second = run_pipeline(stages=["1", "3"], passive_only=True, resume=True)
    assert all(r["status"] == "completed" for r in first.values())
    assert all(r["status"] == "skipped" for r in second.values())
    config.data["targets"]["brands"] = ["ChangedFixture"]
    third = run_pipeline(stages=["1"], passive_only=True, resume=True)
    assert third["1"]["status"] == "completed"
